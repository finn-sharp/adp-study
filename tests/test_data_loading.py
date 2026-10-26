"""다중 형식 로딩·명시적 스키마·기존 UCI/분석 연결을 오프라인 검증합니다."""

from io import BytesIO, StringIO
from dataclasses import fields
import ast
import inspect
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from pandas.testing import assert_frame_equal
from src.datasets import read_table, run_tabular_loader, TabularLoaderProfile
from src.datasets.uci import load_tabular
from src.Recipes import load_profile


class DataLoadingTest(unittest.TestCase):
    def read(self, content, fmt="json", **options):
        source = StringIO(content) if isinstance(content, str) else content
        return read_table(source, overrides=dict(format=fmt, **options))

    def test_same_module_and_field_specification(self):
        from src.datasets import TabularLoaderOutputDTO
        import src.datasets.tabular as module

        tree = ast.parse(Path(module.__file__).read_text())
        for section in range(1, 8):
            self.assertIn("# %d. " % section, Path(module.__file__).read_text())
        for cls in (TabularLoaderProfile, TabularLoaderOutputDTO):
            doc = inspect.getdoc(cls)
            for item in fields(cls):
                self.assertIn(item.name + " (", doc)
        for name in ("read_table", "run_tabular_loader"):
            doc = inspect.getdoc(getattr(module, name))
            for label in ("Args:", "Returns:", "Raises:"):
                self.assertIn(label, doc)
        # Public data classes, not aliases hidden in helper modules.
        self.assertTrue(
            any(
                isinstance(n, ast.ClassDef) and n.name == "TabularLoaderProfile"
                for n in tree.body
            )
        )

    def test_paths_aliases_and_explicit_format(self):
        with tempfile.TemporaryDirectory() as folder:
            for suffix in (".JSON", ".ndjson"):
                path = Path(folder) / ("data" + suffix)
                path.write_text('{"x":1}', encoding="utf-8")
                self.assertEqual(read_table(path).x.tolist(), [1])
            path = Path(folder) / "unknown.bin"
            path.write_bytes(b'[{"x":2}]')
            with self.assertRaises(ValueError):
                read_table(path)
            self.assertEqual(
                read_table(path, overrides={"format": "json"}).x.tolist(), [2]
            )
            with self.assertRaises(FileNotFoundError):
                read_table(Path(folder) / "missing.json")
        with self.assertRaises(ValueError):
            read_table("https://example.com/data.json")
        with self.assertRaises(TypeError):
            read_table(123)
        with self.assertRaises(ValueError):
            read_table(b'[{"x":1}]')

    def test_encoding_and_bom(self):
        data = '[{"이름":"자료"}]'
        self.assertEqual(self.read(data.encode("utf-8-sig")).iloc[0, 0], "자료")
        self.assertEqual(
            self.read(data.encode("cp949"), encoding="cp949").iloc[0, 0], "자료"
        )
        with self.assertRaises(UnicodeError):
            self.read(data.encode("cp949"))

    def test_stream_position_and_ownership(self):
        for stream in (StringIO('XX[{"x":1}]'), BytesIO(b'XX[{"x":1}]')):
            stream.seek(2)
            result = run_tabular_loader(stream, overrides={"format": "json"})
            self.assertEqual(stream.tell(), 2)
            self.assertFalse(stream.closed)
            self.assertTrue(result.metadata["stream_position_restored"])
        stream = StringIO("not JSON")
        with self.assertRaises(ValueError):
            self.read(stream.getvalue())
        with self.assertRaises(ValueError):
            read_table(stream, overrides={"format": "json"})
        self.assertEqual(stream.tell(), 0)

    def test_nonseekable_stream(self):
        class Stream:
            def read(self):
                return '{"x":1}'

        result = run_tabular_loader(Stream(), overrides={"format": "json"})
        self.assertFalse(result.metadata["stream_position_restored"])
        self.assertTrue(result.warnings)

    def test_recipe_profile_and_dto_contract(self):
        profile = load_profile("tabular-loader")
        self.assertIsInstance(profile, TabularLoaderProfile)
        profile.format = "json"
        result = run_tabular_loader(b'[{"x":1},{"x":null}]', profile)
        self.assertEqual(result.n_rows, 2)
        self.assertEqual(result.columns, ["x"])
        self.assertEqual(json.loads(result.to_json())["data"]["data"], [[1.0], [None]])
        with self.assertRaises(ValueError):
            run_tabular_loader(b"{}", profile, overrides={})
        with self.assertRaises(ValueError):
            self.read("{}", typo=1)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "recipe.json"
            path.write_text('{"format":"json","required_columns":["x"]}')
            self.assertEqual(read_table(b'{"x":1}', recipe_path=path).x.tolist(), [1])

    def test_records_paths_and_nested_values(self):
        text = '{"payload":{"rows":[{"x":{"a":1},"items":[2,3]}]}}'
        frame = self.read(text, record_path=["payload", "rows"])
        self.assertEqual(frame.x[0], {"a": 1})
        flat = self.read(text, record_path=["payload", "rows"], flatten=True)
        self.assertEqual(flat["x.a"][0], 1)
        self.assertEqual(flat["items"][0], [2, 3])
        for text, opts in [
            ('{"rows":[]}', {"record_path": ["absent"]}),
            ('[{"a.b":1,"a":{"b":2}}]', {"flatten": True}),
        ]:
            with self.assertRaises(ValueError):
                self.read(text, **opts)

    def test_structured_orientations(self):
        expected = pd.DataFrame({"x": [1, 2], "y": [3, 4]})
        for text, orient in [
            ('{"x":[1,2],"y":[3,4]}', "columns"),
            ('{"a":{"x":1,"y":3},"b":{"x":2,"y":4}}', "index"),
            ('{"columns":["x","y"],"data":[[1,3],[2,4]]}', "split"),
        ]:
            assert_frame_equal(
                self.read(text, orient=orient).reset_index(drop=True), expected
            )
        with self.assertRaises(ValueError):
            self.read('{"columns":["x"],"data":[[1,2]]}', orient="split")

    def test_invalid_json_and_duplicate_columns(self):
        for content in ('{"x":1,"x":2}', '{"x":NaN}', "42", "[1,2]", "{bad"):
            with self.subTest(content=content), self.assertRaises(ValueError):
                self.read(content)
        with self.assertRaises(ValueError):
            self.read('{"columns":["x","x"],"data":[[1,2]]}', orient="split")

    def test_json_lines(self):
        self.assertEqual(self.read('{"x":1}\n\n{"x":2}\n', "jsonl").x.tolist(), [1, 2])
        with self.assertRaisesRegex(ValueError, "2행"):
            self.read('{"x":1}\n{bad}', "jsonl")
        with self.assertRaises(ValueError):
            self.read("[]", "jsonl")
        with self.assertRaises(ValueError):
            self.read('{"x":1}', "jsonl", record_path=["x"])

    def test_txt_lines_preserve_blank_and_literal_na(self):
        frame = self.read("first\n\nNA\n  last  \n", "txt")
        self.assertEqual(frame.text.tolist(), ["first", "", "NA", "  last  "])
        filtered = self.read("a\n\nb\n", "txt", keep_blank_lines=False)
        self.assertEqual(filtered.line_number.tolist(), [1, 3])

    def test_delimiters_quotes_headers_and_blanks(self):
        for content, fmt, options in [
            ("x,y\n1,2\n3,4\n", "csv", {}),
            ("x\ty\n1\t2\n3\t4", "tsv", {}),
            ("x|y\n1|2\n3|4", "txt", {"text_mode": "delimited", "sep": "|"}),
            ("x y\n1 2\n3 4", "txt", {"text_mode": "delimited", "sep": r"\s+"}),
        ]:
            assert_frame_equal(
                self.read(content, fmt, **options),
                pd.DataFrame({"x": [1, 3], "y": [2, 4]}),
            )
        self.assertEqual(
            self.read('x,y\n"a,b","two\nlines"', "csv").iloc[0].tolist(),
            ["a,b", "two\nlines"],
        )
        frame = self.read("x,y\n1,2\n\n3,4", "csv")
        self.assertEqual(len(frame), 3)
        self.assertTrue(frame.iloc[1].isna().all())
        self.assertEqual(
            len(self.read("x,y\n1,2\n\n3,4", "csv", keep_blank_lines=False)), 2
        )
        self.assertEqual(
            self.read("1,2\n3,4", "csv", header=None).columns.tolist(),
            ["column_0", "column_1"],
        )
        self.assertEqual(
            self.read("a,b\n1,2", "csv", names=["x", "y"]).columns.tolist(), ["x", "y"]
        )

    def test_delimited_rejects_data_loss(self):
        for content, options in [
            ("x,x\n1,2", {}),
            ("x,y\n1,2,3", {}),
            ("\nx,y\n1,2", {}),
            ('x,y\n"unterminated,2', {}),
            ("x,y\n1,2", {"header": 0.0}),
        ]:
            with self.subTest(content=content, options=options), self.assertRaises(
                ValueError
            ):
                self.read(content, "csv", **options)
        with self.assertRaises(ValueError):
            self.read("x|y\n1|2", "txt", text_mode="delimited")

    def test_schema_dates_and_missing_rows(self):
        frame = self.read(
            '[{"x":"1","day":"2025-01-02"},{"x":null,"day":null}]',
            dtype={"x": "Float64"},
            parse_dates=["day"],
            date_format="%Y-%m-%d",
            required_columns=["x"],
        )
        self.assertEqual(len(frame), 2)
        self.assertEqual(str(frame.x.dtype), "Float64")
        integers = self.read('[{"x":"1"},{"x":null}]', dtype={"x": "Int64"})
        self.assertEqual(str(integers.x.dtype), "Int64")
        self.assertTrue(pd.isna(integers.x.iloc[1]))
        with self.assertRaises(ValueError):
            self.read('{"x":"invalid"}', dtype={"x": "Float64"})
        self.assertTrue(pd.isna(frame.day.iloc[1]))
        for options in (
            {"required_columns": ["missing"]},
            {"dtype": {"missing": "int64"}},
            {"parse_dates": ["missing"]},
            {"dtype": {"x": "int64"}, "parse_dates": ["x"]},
            {"parse_dates": ["x"]},
            {"date_format": "%Y"},
        ):
            with self.subTest(options=options), self.assertRaises(ValueError):
                self.read('{"x":"not a date"}', **options)

    def test_empty_and_inapplicable_options(self):
        for fmt, content in [("json", "[]"), ("csv", ""), ("txt", "")]:
            with self.assertRaises(ValueError):
                self.read(content, fmt)
            self.assertTrue(self.read(content, fmt, allow_empty=True).empty)
        for options in (
            {"sheet_name": 1},
            {"sep": ","},
            {"xml_xpath": ".//row"},
            {"header": None},
            {"text_mode": "delimited"},
            {"keep_blank_lines": False},
        ):
            with self.subTest(options=options), self.assertRaises(ValueError):
                self.read('{"x":1}', **options)

    def test_optional_dependency_error(self):
        with patch("importlib.import_module", side_effect=ImportError("missing")):
            from src.datasets._readers import optional_package

            with self.assertRaisesRegex(ImportError, "requirements-data-loading.txt"):
                optional_package("yaml", "PyYAML")

    @unittest.skipUnless(importlib.util.find_spec("yaml"), "optional PyYAML")
    def test_yaml_and_rejected_structures(self):
        assert_frame_equal(
            self.read("rows:\n  - x: 1\n  - x: 2", "yaml", record_path=["rows"]),
            pd.DataFrame({"x": [1, 2]}),
        )
        for content in (
            "x: 1\nx: 2",
            "x: !!python/object:os.system {}",
            "x: &a [*a]",
            "1: x",
            "x: !!set {a: null}",
            "x: {<<: {a: 1}}",
        ):
            with self.subTest(content=content), self.assertRaises(ValueError):
                self.read(content, "yaml")
        frame = self.read("x: 2025-01-02", "yaml")
        result = run_tabular_loader(
            StringIO("x: 2025-01-02"), overrides={"format": "yaml"}
        )
        self.assertIn("2025-01-02", result.to_json())
        self.assertEqual(len(frame), 1)

    @unittest.skipUnless(importlib.util.find_spec("defusedxml"), "optional defusedxml")
    def test_xml_paths_attributes_and_namespace(self):
        text = '<root xmlns:n="urn:test"><n:row id="a"><n:x unit="m">1</n:x></n:row></root>'
        frame = self.read(
            text, "xml", xml_xpath="./n:row", xml_namespaces={"n": "urn:test"}
        )
        self.assertEqual(
            frame.iloc[0].to_dict(),
            {"@id": "a", "{urn:test}x": "1", "{urn:test}x@unit": "m"},
        )
        assert_frame_equal(
            self.read(
                "<root><row><x>1</x></row><row><x>2</x></row></root>",
                "xml",
                dtype={"x": "int64"},
            ),
            pd.DataFrame({"x": [1, 2]}),
        )
        self.assertTrue(self.read("<root/>", "xml", allow_empty=True).empty)

    @unittest.skipUnless(importlib.util.find_spec("defusedxml"), "optional defusedxml")
    def test_xml_rejects_unsafe_or_lossy_structure(self):
        for content in (
            '<!DOCTYPE root [<!ENTITY x "bad">]><root><row>&x;</row></root>',
            "<root><row><x>1</x><x>2</x></row></root>",
            "<root><row><x><a>1</a></x></row></root>",
            "<root><row>text<x>1</x></row></root>",
            "<root>",
        ):
            with self.subTest(content=content), self.assertRaises(ValueError):
                self.read(content, "xml")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional openpyxl")
    def test_xlsx_sheets_formulas_and_headers(self):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "자료"
        sheet.append(["x", "y"])
        sheet.append([1, 2])
        sheet.append([None, None])
        sheet.append([3, "=1+1"])
        other = book.create_sheet("다른표")
        other.append([7, 8])
        stream = BytesIO()
        book.save(stream)
        book.close()
        payload = stream.getvalue()
        frame = self.read(payload, "xlsx", sheet_name="자료")
        self.assertEqual(len(frame), 3)
        self.assertTrue(pd.isna(frame.iloc[-1, 1]))
        self.assertEqual(
            self.read(payload, "xlsx", excel_data_only=False).iloc[-1, 1], "=1+1"
        )
        self.assertEqual(len(self.read(payload, "xlsx", keep_blank_lines=False)), 2)
        self.assertEqual(
            self.read(payload, "xlsx", sheet_name=1, header=None).iloc[0].tolist(),
            [7, 8],
        )
        for name in (5, "missing"):
            with self.assertRaises(ValueError):
                self.read(payload, "xlsx", sheet_name=name)
        with self.assertRaises(ValueError):
            self.read(b"not an excel file", "xlsx")
        sheet = Workbook()
        sheet.active.append(["x", "x"])
        sheet.active.append([1, 2])
        stream = BytesIO()
        sheet.save(stream)
        sheet.close()
        with self.assertRaises(ValueError):
            self.read(stream.getvalue(), "xlsx")

    def test_custom_target_adapter_and_validation(self):
        text = b'[{"x":1,"outcome":0},{"x":2,"outcome":1}]'
        data = load_tabular(
            "classification",
            binary=True,
            source=text,
            target_column="outcome",
            data_overrides={"format": "json"},
        )
        self.assertEqual(data.columns.tolist(), ["x", "target"])
        for content in (b'{"x":1}', b'{"target":null,"x":1}', b'{"target":1}'):
            with self.assertRaises(ValueError):
                load_tabular(source=content, data_overrides={"format": "json"})
        with self.assertRaises(ValueError):
            load_tabular(
                source=b'{"target":1,"outcome":2,"x":3}',
                target_column="outcome",
                data_overrides={"format": "json"},
            )
        with self.assertRaises(ValueError):
            load_tabular(
                source=text, cache_dir="unused", data_overrides={"format": "json"}
            )

    def test_existing_uci_path_is_preserved(self):
        # 실제 다운로드 대신 기존 원본 CSV 응답을 대입하여 선택·타깃 계약을 확인합니다.
        wine = b"fixed acidity;quality\n7;5\n8;6\n"
        with patch("src.datasets.uci.dataset_bytes", return_value=wine):
            frame = load_tabular("regression")
        assert_frame_equal(
            frame, pd.DataFrame({"fixed acidity": [7, 8], "target": [5, 6]})
        )

    def test_linear_skeleton_accepts_custom_file(self):
        from src.MachineLearning.skeleton.linear_regression import main

        rows = [{"x": i, "outcome": 2 * i + 3} for i in range(40)]
        result = main(
            source=json.dumps(rows).encode(),
            target_column="outcome",
            data_overrides={"format": "json"},
        )
        self.assertEqual(result.dataset_name, "사용자 파일")
        self.assertEqual(result.n_rows, 40)
        self.assertGreater(result.metrics["r2"], 0.99)


if __name__ == "__main__":
    unittest.main()
